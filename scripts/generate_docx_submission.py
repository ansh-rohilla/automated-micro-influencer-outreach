"""
Generate a Microsoft Word (.docx) Submission Document for the project.
Includes all 8 submission items from the specification with screenshot placeholders and tables.
"""

import os
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn


def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding in dxa."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_callout_box(doc, text_paragraphs, border_color="2563EB", bg_color="F8FAFC"):
    """Creates a callout box with a colored left accent border."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    # Set borders: left thick, others none
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

    cell_p = cell.paragraphs[0]
    for i, line in enumerate(text_paragraphs):
        if i > 0:
            cell_p = cell.add_paragraph()
        cell_p.paragraph_format.space_before = Pt(2)
        cell_p.paragraph_format.space_after = Pt(2)
        cell_p.paragraph_format.line_spacing = 1.15
        
        # Check if line has bold prefix
        if ":" in line and not line.startswith("http"):
            parts = line.split(":", 1)
            r1 = cell_p.add_run(parts[0] + ":")
            r1.bold = True
            r1.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            r2 = cell_p.add_run(parts[1])
            r2.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        else:
            r = cell_p.add_run(line)
            r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)


def add_screenshot_placeholder(doc, label, description):
    """Adds a visual placeholder box for screenshots."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=300, bottom=300, left=250, right=250)

    # Dashed border
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="dashed" w:sz="12" w:space="0" w:color="94A3B8"/>
            <w:left w:val="dashed" w:sz="12" w:space="0" w:color="94A3B8"/>
            <w:bottom w:val="dashed" w:sz="12" w:space="0" w:color="94A3B8"/>
            <w:right w:val="dashed" w:sz="12" w:space="0" w:color="94A3B8"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_icon = p.add_run("📷 [IMAGE PLACEHOLDER]\n")
    r_icon.bold = True
    r_icon.font.size = Pt(12)
    r_icon.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    r_label = p.add_run(f"{label}\n")
    r_label.bold = True
    r_label.font.size = Pt(11)
    r_label.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    r_desc = p.add_run(f"({description})")
    r_desc.italic = True
    r_desc.font.size = Pt(9.5)
    r_desc.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # Space after table
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(6)


def create_submission_docx(output_path="submission_document.docx"):
    doc = Document()

    # Page Margins (1 inch all around)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Style: Default Font Calibri
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # Document Header / Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    title_run = title_p.add_run("Automated Micro-Influencer Outreach System")
    title_run.bold = True
    title_run.font.size = Pt(22)
    title_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_before = Pt(0)
    subtitle_p.paragraph_format.space_after = Pt(14)
    sub_run = subtitle_p.add_run("Formal Project Submission Report & Technical Documentation")
    sub_run.font.size = Pt(13)
    sub_run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    # Meta Info Table
    meta_table = doc.add_table(rows=3, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_info = [
        ("Candidate Name:", "Ansh Rohilla", "Target Niche:", "Technology & AI"),
        ("GitHub Repository:", "https://github.com/ansh-rohilla/automated-micro-influencer-outreach", "Test Run Yield:", "58 Real Creators (30 Micro Shortlisted)"),
        ("Sending Modes:", "Simulation + Live SMTP", "Evaluation Status:", "Fully Verified & Passing (6/6 Tests)")
    ]
    for r_idx, (k1, v1, k2, v2) in enumerate(meta_info):
        row = meta_table.rows[r_idx]
        c1, c2 = row.cells[0], row.cells[1]
        set_cell_background(c1, "F8FAFC")
        set_cell_background(c2, "F8FAFC")
        set_cell_margins(c1, 60, 60, 100, 100)
        set_cell_margins(c2, 60, 60, 100, 100)

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        r_k1 = p1.add_run(k1 + " ")
        r_k1.bold = True
        p1.add_run(v1)

        p2 = c2.paragraphs[0]
        p2.paragraph_format.space_before = Pt(2)
        p2.paragraph_format.space_after = Pt(2)
        r_k2 = p2.add_run(k2 + " ")
        r_k2.bold = True
        p2.add_run(v2)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ==========================================
    # ITEM 1: GITHUB REPOSITORY & PROJECT FILES
    # ==========================================
    h1 = doc.add_heading("1. GitHub Repository or Project Files", level=1)
    h1.runs[0].font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    p_repo = doc.add_paragraph()
    p_repo.add_run("The complete codebase, datasets, test suites, and interactive web dashboard are version-controlled and publicly hosted on GitHub at:\n")
    r_link = p_repo.add_run("https://github.com/ansh-rohilla/automated-micro-influencer-outreach")
    r_link.bold = True
    r_link.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    doc.add_paragraph(
        "The project is structured into modular packages following industry design patterns:"
    )

    proj_tree = [
        "• main.py — Master one-command end-to-end pipeline runner",
        "• app.py — Interactive Streamlit web application dashboard",
        "• src/discovery/ — Public directory web scraping & multi-platform extraction",
        "• src/filtering/ — Deterministic qualification rules & brand-fit scoring engine",
        "• src/enrichment/ — Contact discovery, audience demographics & tone inference",
        "• src/personalization/ — Dynamic prompt engineering & strict word-count validator",
        "• src/sending/ — Dual-mode dispatcher (SMTP/Simulation) & duplicate outreach prevention",
        "• tests/test_pipeline.py — Comprehensive unit tests covering all components",
        "• data/ — Raw and processed CSV / JSON datasets matching all specification formats"
    ]
    for item in proj_tree:
        p_item = doc.add_paragraph(item)
        p_item.paragraph_format.space_before = Pt(1)
        p_item.paragraph_format.space_after = Pt(2)

    # ==========================================
    # ITEM 2: README / DOCUMENTATION
    # ==========================================
    h2 = doc.add_heading("2. README / Documentation Summary", level=1)
    h2.runs[0].font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    doc.add_paragraph(
        "The project repository contains an exhaustive README.md and a formal SUBMISSION.md addressing all 11 required technical aspects:"
    )

    doc_points = [
        ("Technology Stack: ", "Python 3.9+, Pydantic v2, BeautifulSoup4, Pandas, LiteLLM, smtplib, Streamlit."),
        ("Data Sources & Discovery: ", "Scraped live from public creator directories (Collabstr Tech & AI categories) yielding 58 verified creators."),
        ("Filtering Logic: ", "Strict 5,000 to 100,000 follower bounds, ≥2.0% engagement rate, brand safety negative filters, and a 0–10 brand-fit scoring heuristic."),
        ("Profile Enrichment: ", "Mandatory and optional fields populated under strict anti-fabrication standards (real verified business emails or explicit 'Not Found')."),
        ("AI Personalization: ", "Dynamic prompt templates generating bespoke 60–90 word email pitches and 15–30 word Instagram DMs."),
        ("Sending Mechanism: ", "Safe simulation mode + live SMTP authenticated dispatch with duplicate email detection (idempotency).")
    ]
    for bold_text, desc_text in doc_points:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        r_b = p.add_run("• " + bold_text)
        r_b.bold = True
        p.add_run(desc_text)

    # ==========================================
    # ITEM 3: WORKING DEMO OR SCREENSHOTS / VIDEO
    # ==========================================
    h3 = doc.add_heading("3. Working Demo or Screenshots / Video", level=1)
    h3.runs[0].font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    doc.add_paragraph(
        "The system includes a production-grade Streamlit web dashboard. Evaluators can launch the interactive interface locally using:"
    )
    p_code = doc.add_paragraph()
    r_c = p_code.add_run("    streamlit run app.py")
    r_c.font.name = "Consolas"
    r_c.bold = True
    r_c.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    doc.add_paragraph(
        "Below are placeholders to attach your walkthrough screenshots or video link:"
    )

    add_screenshot_placeholder(
        doc,
        label="Screenshot 1: Dashboard Overview & Pipeline KPIs",
        description="Insert screenshot of Dashboard Overview showing 58 Discovered, 30 Shortlisted, 28 Excluded, and System Architecture"
    )

    add_screenshot_placeholder(
        doc,
        label="Screenshot 2: Influencer Discovery Explorer",
        description="Insert screenshot of 1. Influencer Discovery tab displaying the interactive table of 58 real Tech & AI creators"
    )

    add_screenshot_placeholder(
        doc,
        label="Screenshot 3: Filtering & Classification Audit Table",
        description="Insert screenshot of 2. Filtering & Classification tab showing PASSED vs FAILED badges and explicit rejection reasons"
    )

    add_screenshot_placeholder(
        doc,
        label="Screenshot 4: AI Message Personalization & Word Count Validator",
        description="Insert screenshot of 4. AI Message Personalization tab showing the 60-90w Email Pitch and 15-30w Instagram DM"
    )

    add_screenshot_placeholder(
        doc,
        label="Screenshot 5: Sending Layer & Idempotent Outreach Tracker",
        description="Insert screenshot of 5. Sending Layer & Tracker showing SIMULATED_SENT, SKIPPED_NO_EMAIL, and SKIPPED_DUPLICATE"
    )

    # ==========================================
    # ITEM 4: INFLUENCER DATASET
    # ==========================================
    h4 = doc.add_heading("4. Influencer Dataset (Page 7 Specification)", level=1)
    h4.runs[0].font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    doc.add_paragraph(
        "The system retrieved 58 verified real creators during the test run (exceeding the requirement of at least 50). "
        "The complete dataset is stored in data/influencer_dataset.csv matching the exact recommended format:"
    )

    # Read consolidated dataset
    df_data = pd.read_csv("data/influencer_dataset.csv")
    sample_df = df_data.head(10)

    tbl_data = doc.add_table(rows=len(sample_df) + 1, cols=7)
    tbl_data.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Name", "Platform", "Followers", "Engagement", "Niche", "Contact Email", "Status"]

    for col_idx, h_text in enumerate(headers):
        cell = tbl_data.cell(0, col_idx)
        cell.paragraphs[0].add_run(h_text).bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_background(cell, "E2E8F0")
        set_cell_margins(cell, 80, 80, 80, 80)

    for row_idx, (_, row) in enumerate(sample_df.iterrows(), start=1):
        r_cells = tbl_data.rows[row_idx].cells
        vals = [
            str(row["Name"]),
            str(row["Platform"]),
            f"{int(row['Followers']):,}",
            str(row["Engagement"]),
            str(row["Niche"]),
            str(row["Email"]),
            str(row["Status"])
        ]
        bg = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(vals):
            cell = r_cells[c_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            if c_idx == 6:  # Status
                r.bold = True
                if val == "PASSED":
                    r.font.color.rgb = RGBColor(0x06, 0x5F, 0x46)
                else:
                    r.font.color.rgb = RGBColor(0x99, 0x1B, 0x1B)
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 80, 80)

    p_cap = doc.add_paragraph()
    p_cap.paragraph_format.space_before = Pt(4)
    r_cap = p_cap.add_run("Table 1: Sample from 58-creator dataset. See data/influencer_dataset.csv for all 58 rows.")
    r_cap.font.size = Pt(9)
    r_cap.italic = True
    r_cap.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # ==========================================
    # ITEM 5: SAMPLE PERSONALIZED OUTREACH MESSAGES
    # ==========================================
    h5 = doc.add_heading("5. Sample Personalized Outreach Messages", level=1)
    h5.runs[0].font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    doc.add_paragraph(
        "Each shortlisted micro-influencer receives two distinct personalized messages. "
        "Lengths are strictly enforced (Email: 60–90 words | Instagram DM: 15–30 words) referencing the creator's "
        "tone, audience, recent content topics, and assigned collaboration angle."
    )

    # Sample 1
    doc.add_heading("Creator 1: Diego Montes B (@diegoinnovacion) — TikTok (56,800 followers)", level=2)
    add_callout_box(
        doc,
        [
            "Collaboration Angle: Sponsorship",
            "Value Proposition: Competitive flat creator fees with creative freedom on campaign integration.",
            "Email Collaboration Pitch (74 words | Valid: 60-90w):",
            "Subject: Collaboration inquiry with Diego – Sponsorship for Practical Generative AI tools and workflow automation",
            "Hi Diego, I loved your recent breakdown on practical generative ai tools and workflow automation. Your professional, business-focused, strategic insights and case studies and highly engaged 18-34 tech audience in Lima align seamlessly with our team's mission. We are launching an upcoming sponsorship and would love to partner with you. We offer competitive flat creator fees with creative freedom on campaign integration. Would you be open to reviewing a brief partnership brief this week?",
            "Instagram Direct Message (28 words | Valid: 15-30w):",
            "\"Hi Diego, loved your recent practical generative ai tools and workflow automation content! Your tech audience looks like a perfect fit for our upcoming sponsorship. Open to collaborating?\""
        ],
        border_color="2563EB",
        bg_color="F8FAFC"
    )

    # Sample 2
    doc.add_heading("Creator 2: Kimberlyn Padilla (@kimberlynpadilla) — TikTok (60,200 followers)", level=2)
    add_callout_box(
        doc,
        [
            "Collaboration Angle: UGC Content Creation",
            "Value Proposition: Paid commercial licensing deal for our upcoming high-growth ad campaigns.",
            "Email Collaboration Pitch (77 words | Valid: 60-90w):",
            "Subject: Collaboration inquiry with Kimberlyn – UGC content creation for Practical Generative AI tools and workflow automation",
            "Hi Kimberlyn, I loved your recent breakdown on practical generative ai tools and workflow automation. Your engaging, practical product demos and relatable consumer reviews and highly engaged 18-34 tech audience in Cucuta align seamlessly with our team's mission. We are launching an upcoming ugc content creation and would love to partner with you. We offer paid commercial licensing deal for our upcoming high-growth ad campaigns. Would you be open to reviewing a brief partnership brief this week?",
            "Instagram Direct Message (30 words | Valid: 15-30w):",
            "\"Hi Kimberlyn, loved your recent practical generative ai tools and workflow automation content! Your tech audience looks like a perfect fit for our upcoming ugc content creation. Open to collaborating?\""
        ],
        border_color="A855F7",
        bg_color="FAF5FF"
    )

    # Sample 3
    doc.add_heading("Creator 3: Abdelhamid Oughanem (@abdelhamidhania) — Instagram (13,600 followers)", level=2)
    add_callout_box(
        doc,
        [
            "Collaboration Angle: Paid Product Placement",
            "Value Proposition: Dedicated sponsored demo segment highlighting your authentic workflow.",
            "Email Collaboration Pitch (74 words | Valid: 60-90w):",
            "Subject: Collaboration inquiry with Abdelhamid – Paid product placement for Practical Generative AI tools and workflow automation",
            "Hi Abdelhamid, I loved your recent breakdown on practical generative ai tools and workflow automation. Your technical, educational, in-depth tutorials and code walkthroughs and highly engaged 18-34 tech audience in Ottawa align seamlessly with our team's mission. We are launching an upcoming paid product placement and would love to partner with you. We offer dedicated sponsored demo segment highlighting your authentic workflow. Would you be open to reviewing a brief partnership brief this week?",
            "Instagram Direct Message (30 words | Valid: 15-30w):",
            "\"Hi Abdelhamid, loved your recent practical generative ai tools and workflow automation content! Your tech audience looks like a perfect fit for our upcoming paid product placement. Open to collaborating?\""
        ],
        border_color="059669",
        bg_color="F0FDF4"
    )

    # ==========================================
    # ITEM 6: AUTOMATION WORKFLOW & OUTREACH TRACKER
    # ==========================================
    h6 = doc.add_heading("6. Automation Workflow & Outreach Tracker", level=1)
    h6.runs[0].font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    doc.add_paragraph(
        "The sending layer implements an automated delivery and auditing workflow: "
        "1. Selects influencers with valid emails; 2. Verifies duplicate prevention registry (idempotency); "
        "3. Delivers via configured mode (Simulation or live SMTP); 4. Records sending status and timestamps in data/processed/outreach_tracker.csv."
    )

    tracker_df = pd.read_csv("data/processed/outreach_tracker.csv").head(6)
    tbl_tracker = doc.add_table(rows=len(tracker_df) + 1, cols=5)
    tbl_tracker.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_headers = ["Influencer", "Email", "Message Preview", "Sent Date", "Status"]

    for col_idx, h_text in enumerate(t_headers):
        cell = tbl_tracker.cell(0, col_idx)
        cell.paragraphs[0].add_run(h_text).bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_background(cell, "E2E8F0")
        set_cell_margins(cell, 80, 80, 80, 80)

    for row_idx, (_, row) in enumerate(tracker_df.iterrows(), start=1):
        r_cells = tbl_tracker.rows[row_idx].cells
        msg_prev = str(row["Message_Generated"])[:45] + "..."
        vals = [
            str(row["Influencer"]),
            str(row["Email"]),
            msg_prev,
            str(row["Sent_Date"]).split("T")[0],
            str(row["Status"])
        ]
        bg = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(vals):
            cell = r_cells[c_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            if c_idx == 4:
                r.bold = True
                if "SENT" in val:
                    r.font.color.rgb = RGBColor(0x06, 0x5F, 0x46)
                elif "DUPLICATE" in val:
                    r.font.color.rgb = RGBColor(0x92, 0x40, 0x0E)
                else:
                    r.font.color.rgb = RGBColor(0x03, 0x69, 0xA1)
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 80, 80)

    # ==========================================
    # ITEM 7: SETUP INSTRUCTIONS
    # ==========================================
    h7 = doc.add_heading("7. Setup Instructions", level=1)
    h7.runs[0].font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    doc.add_paragraph("Follow these steps to run and test the complete pipeline locally:")
    steps = [
        "1. Clone repository:\n   git clone https://github.com/ansh-rohilla/automated-micro-influencer-outreach.git\n   cd automated-micro-influencer-outreach",
        "2. Create virtual environment & install dependencies:\n   python3 -m venv venv\n   source venv/bin/activate\n   pip install -r requirements.txt",
        "3. Run entire end-to-end pipeline:\n   python main.py",
        "4. Launch interactive Streamlit web dashboard:\n   streamlit run app.py",
        "5. Execute test suite (6/6 passing):\n   python -m unittest discover tests"
    ]
    for s in steps:
        p_s = doc.add_paragraph(s)
        p_s.paragraph_format.space_before = Pt(3)
        p_s.paragraph_format.space_after = Pt(4)

    # ==========================================
    # ITEM 8: LIST OF APIS AND TOOLS USED
    # ==========================================
    h8 = doc.add_heading("8. List of APIs and Tools Used", level=1)
    h8.runs[0].font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    tools = [
        ("Python 3.9+ / 3.13: ", "Primary programming language and runtime environment."),
        ("Pydantic v2: ", "Strict schema validation and data modeling for all stages."),
        ("BeautifulSoup4 & urllib: ", "Web scraping and HTML parsing of public creator directories."),
        ("Pandas: ", "Data structuring, deduplication, filtering, and CSV/JSON persistence."),
        ("LiteLLM & Google GenAI SDK: ", "AI message personalization and dynamic prompt generation."),
        ("smtplib & email.mime: ", "Standard RFC-822 email construction, live SMTP, and simulation mode."),
        ("Streamlit: ", "Interactive multi-page web application dashboard."),
        ("python-docx: ", "Automated generation of this Microsoft Word submission report.")
    ]
    for b_tool, desc in tools:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r_b = p.add_run("• " + b_tool)
        r_b.bold = True
        p.add_run(desc)

    doc.save(output_path)
    print(f"Generated Word submission document at: {output_path}")


if __name__ == "__main__":
    create_submission_docx("submission_document.docx")
